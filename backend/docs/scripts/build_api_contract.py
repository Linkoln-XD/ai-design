#!/usr/bin/env python3
"""Собирает docs/API-контракт.docx из этого сценария.

Запуск из корня репозитория:

    python3 -m pip install -r docs/requirements.txt
    python3 docs/scripts/build_api_contract.py
"""

from __future__ import annotations

import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit(
        "Не установлен python-docx. Выполните: "
        "python3 -m pip install -r docs/requirements.txt"
    )


OUT = Path(__file__).resolve().parents[1] / "API-контракт.docx"
NAVY = RGBColor(0x1F, 0x4E, 0x79)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
CODE_BG = "F4F6F8"
HEADER_BG = "1F4E79"
NOTE_BG = "EEF4FA"


def _rfonts(rpr, name: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _lang(rpr) -> None:
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "ru-RU")
    lang.set(qn("w:eastAsia"), "ru-RU")


def style_font(style, name: str, size: int, bold: bool, color: RGBColor | None) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    _rfonts(rpr, name)
    _lang(rpr)


def shade_cell(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def set_run_font(run, name: str, size: int, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.name = name
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    _rfonts(rpr, name)
    _lang(rpr)


def prevent_row_split(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trpr.append(cant)


def mark_header_row(row) -> None:
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    trpr.append(header)


def set_table_width(table) -> None:
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "pct")
    tblw.set(qn("w:w"), "5000")


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    set_run_font(run, "Arial", 9, color=NAVY)


class Spec:
    def __init__(self) -> None:
        self.doc = Document()
        self._setup()

    def _setup(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.header_distance = Cm(0.6)
        section.footer_distance = Cm(0.6)

        normal = self.doc.styles["Normal"]
        style_font(normal, "Arial", 11, False, None)
        normal.paragraph_format.space_after = Pt(8)
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.line_spacing = 1.12

        for level, size, color in (
            (1, 16, NAVY),
            (2, 13, BLUE),
            (3, 12, NAVY),
        ):
            style = self.doc.styles[f"Heading {level}"]
            style_font(style, "Arial", size, True, color)
            style.paragraph_format.space_before = Pt(16 if level == 1 else 12)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.line_spacing = 1.1

        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = header.add_run("Arachis Assistant  ·  контракт HTTP API  ·  версия 1.0")
        set_run_font(run, "Arial", 9, color=NAVY)

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
        left = footer.add_run("Backend, веб-клиент и модуль ассистента")
        set_run_font(left, "Arial", 9, color=NAVY)
        footer.add_run("          ")
        add_page_number(footer)

        core = self.doc.core_properties
        core.title = "Arachis Assistant — контракт HTTP API, версия 1.0"
        core.subject = "Контракт API календаря задач, встреч и напоминаний"
        core.category = "Спецификация"
        core.language = "ru-RU"

    def h(self, text: str, level: int = 1):
        return self.doc.add_heading(text, level=level)

    def p(self, text: str):
        paragraph = self.doc.add_paragraph()
        run = paragraph.add_run(text)
        set_run_font(run, "Arial", 11)
        return paragraph

    def bullet(self, text: str) -> None:
        paragraph = self.doc.add_paragraph(style="List Bullet")
        paragraph.clear()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(0)
        run = paragraph.add_run(text)
        set_run_font(run, "Arial", 11)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            self.bullet(item)

    def note(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        set_table_width(table)
        cell = table.cell(0, 0)
        shade_cell(cell, NOTE_BG)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.space_before = Pt(2)
        run = paragraph.add_run(text)
        set_run_font(run, "Arial", 10)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def table(self, headers: list[str], rows: list[list[str]]) -> None:
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        set_table_width(table)
        for index, header in enumerate(headers):
            cell = table.rows[0].cells[index]
            shade_cell(cell, HEADER_BG)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            run = paragraph.add_run(header)
            set_run_font(run, "Arial", 9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        mark_header_row(table.rows[0])
        prevent_row_split(table.rows[0])
        for r_index, row in enumerate(rows):
            for c_index, value in enumerate(row):
                cell = table.rows[r_index + 1].cells[c_index]
                if r_index % 2 == 1:
                    shade_cell(cell, "F7F9FB")
                cell.text = ""
                paragraph = cell.paragraphs[0]
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                run = paragraph.add_run(value)
                set_run_font(run, "Arial", 9)
            prevent_row_split(table.rows[r_index + 1])
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.space_before = Pt(0)

    def code(self, text: str) -> None:
        table = self.doc.add_table(rows=1, cols=1)
        set_table_width(table)
        cell = table.cell(0, 0)
        shade_cell(cell, CODE_BG)
        lines = text.strip("\n").split("\n")
        for index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            run = paragraph.add_run(line if line else " ")
            set_run_font(run, "Courier New", 8)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(6)
        spacer.paragraph_format.space_before = Pt(2)

    def save(self) -> None:
        self.doc.save(OUT)


def build() -> None:
    spec = Spec()
    cover(spec)
    purpose(spec)
    conventions(spec)
    errors(spec)
    auth(spec)
    tasks(spec)
    meetings(spec)
    reminders(spec)
    calendar(spec)
    search(spec)
    notifications(spec)
    history(spec)
    scenarios(spec)
    out_of_scope(spec)
    changelog(spec)
    spec.save()


def cover(s: Spec) -> None:
    title = s.doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    run = title.add_run("Arachis Assistant")
    set_run_font(run, "Arial", 28, bold=True, color=NAVY)

    subtitle = s.doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Контракт HTTP API")
    set_run_font(run, "Arial", 22, bold=True, color=BLUE)

    meta = s.doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Версия 1.0  ·  10 октября 2026")
    set_run_font(run, "Arial", 12, color=NAVY)

    status = s.doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run("Рабочая спецификация для команды")
    set_run_font(run, "Arial", 12)

    s.doc.add_page_break()
    s.h("Содержание", 1)
    for line in (
        "1. Назначение и роли",
        "2. Общие соглашения",
        "3. Ошибки",
        "4. Аутентификация",
        "5. Задачи",
        "6. Встречи",
        "7. Напоминания",
        "8. Списки и календарь",
        "9. Поиск",
        "10. Уведомления",
        "11. Импорт и экспорт истории",
        "12. Сценарии вызовов",
        "13. За границами версии 1.0",
        "14. История версий",
    ):
        s.bullet(line)


def purpose(s: Spec) -> None:
    s.h("1. Назначение и роли", 1)
    s.p(
        "Документ описывает HTTP API личного календаря: задачи, встречи и напоминания. "
        "Это общий контракт трёх частей системы. Пока сервер строится, источником договорённости является этот документ."
    )
    s.table(
        ["Роль", "Ответственность"],
        [
            [
                "Backend",
                "Реализует сервер: проверки, хранение, поиск, уведомления, импорт и экспорт.",
            ],
            [
                "Веб-клиент",
                "Экраны календаря и форм. Вызывает те же методы, что и ассистент.",
            ],
            [
                "Модуль ассистента",
                "Разбирает реплику, переводит относительные даты в абсолютные и вызывает API. Уточняющие вопросы и подтверждение удаления ведёт сам.",
            ],
        ],
    )
    s.p("Сервер принимает уже разобранное действие. Фразы «завтра», «вечером», «в пятницу» и местоимение «её» к серверу не приходят.")
    s.p(
        "Правила продукта, которые клиент применяет до вызова: «утром» — 08:00, «днём» — 12:00, "
        "«вечером» — 18:00, точное время важнее времени суток, недельный обзор считается с понедельника по воскресенье. "
        "Всё это считается в часовом поясе пользователя. Сервер сравнивает уже готовые дату и время с текущим моментом."
    )
    s.h("Слова в этом документе", 2)
    s.table(
        ["Слово", "Значение"],
        [
            ["Запись", "Задача, встреча или напоминание."],
            ["Клиент", "Веб-клиент или модуль ассистента."],
            ["Пользователь", "Владелец личного календаря."],
            ["Должен", "Обязательное поведение сервера или клиента."],
        ],
    )


def conventions(s: Spec) -> None:
    s.h("2. Общие соглашения", 1)
    s.h("Адрес и формат", 2)
    s.p("Базовый путь: https://{host}/api/v1. Конкретный хост сообщается отдельно, когда сервер доступен заказчику. Обмен идёт по HTTPS.")
    s.bullets(
        [
            "Тело запроса и ответа — JSON, кодировка UTF-8, заголовок Content-Type: application/json.",
            "Имена полей — snake_case на английском. Значения перечислений — английские слова в нижнем регистре.",
            "Тексты message — на русском. Клиент ветвится по полю code, а не по тексту message.",
            "В ответах присутствуют все поля объекта. Пустое значение передаётся как null, пустой список — как [].",
            "В запросах пропущенное поле означает «не передано». JSON null означает «очистить», и только там, где это прямо разрешено.",
            "Идентификаторы записей — UUID версии 4, которые назначает сервер. Клиент свой UUID записи не выбирает.",
        ]
    )
    s.h("Время", 2)
    s.p(
        "Момент времени передаётся в формате RFC 3339 с часовым поясом: 2026-10-11T18:00:00+03:00. "
        "Суффикс Z сервер принимает как UTC. Календарная дата передаётся как 2026-10-11."
    )
    s.p(
        "В профиле пользователя хранится пояс IANA, по умолчанию Europe/Moscow. "
        "Сервер сохраняет моменты как абсолютное время и в ответах показывает их в текущем поясе пользователя. "
        "Срок задачи — календарная дата без часа. Её нужно вычислять в поясе пользователя до отправки: вечер пятницы по Москве не должен превращаться в субботу из-за перевода в UTC."
    )
    s.p(
        "Сравнение с прошлым сервер делает по своим часам. Допуск на расхождение часов — 60 секунд. "
        "Дата срока сравнивается с сегодняшней датой в поясе пользователя: сегодняшний день ещё допустим, вчерашний уже нет."
    )
    s.note(
        "Пример. Пояс пользователя Europe/Moscow. Клиент прислал начало встречи 2026-10-11T15:00:00Z. "
        "Сервер сохраняет этот же момент и возвращает 2026-10-11T18:00:00+03:00."
    )
    s.h("Авторизация запросов", 2)
    s.p(
        "Защищённые методы требуют заголовок Authorization: Bearer <access_token>. "
        "Сервер не использует cookie. Чужая запись для вызывающего неотличима от отсутствующей: ответ 404 not_found."
    )
    s.h("Идемпотентность", 2)
    s.p(
        "На создание задачи, встречи и напоминания, а также на удаление, клиент передаёт заголовок Idempotency-Key: произвольная строка от 1 до 128 символов, удобно UUID. "
        "Ключ действует сутки и привязан к пользователю, методу и пути."
    )
    s.bullets(
        [
            "Повтор с тем же ключом и тем же телом возвращает исходный код и исходное тело, новую запись сервер не создаёт.",
            "Тот же ключ с другим телом даёт 409 idempotency_conflict.",
            "Два шага создания встречи при пересечении — это два разных ключа, потому что тела различаются полем allow_overlap.",
            "Без заголовка сервер выполняет запрос как новый.",
        ]
    )
    s.h("Списки", 2)
    s.p("Методы списков принимают limit и offset. По умолчанию limit равен 100, наибольшее значение — 200. offset по умолчанию равен 0.")
    s.code(
        """
{
  "items": [],
  "limit": 100,
  "offset": 0,
  "total": 0
}
"""
    )
    s.h("Браузер", 2)
    s.p(
        "Сервер разрешает обращения с адреса веб-клиента и заголовки Authorization, Content-Type и Idempotency-Key. "
        "Токен доступа веб-клиент хранит у себя и подставляет в заголовок."
    )


def errors(s: Spec) -> None:
    s.h("3. Ошибки", 1)
    s.p("Тело ошибки устроено одинаково. У ответа 204 тела нет.")
    s.code(
        """
{
  "code": "validation_error",
  "message": "Некорректные данные",
  "details": {
    "errors": [
      { "field": "title", "code": "required", "message": "Нужно название" }
    ]
  }
}
"""
    )
    s.p("Сервер возвращает все найденные ошибки полей сразу. Если обязательное поле пустое, проверка «время в прошлом» для него не выполняется.")
    s.table(
        ["HTTP", "code", "Когда"],
        [
            ["401", "unauthorized", "Нет токена или токен не распознан."],
            ["401", "token_expired", "Срок access token истёк. Клиент обновляет его и повторяет запрос один раз."],
            ["401", "invalid_credentials", "Неверная пара почта и пароль. Текст один и тот же, есть такая почта или нет."],
            ["401", "invalid_refresh", "Refresh token неизвестен, просрочен или уже использован."],
            ["404", "not_found", "Записи нет, она другого типа или принадлежит другому пользователю."],
            ["409", "email_taken", "Почта уже зарегистрирована."],
            ["409", "overlap", "Встреча пересекается с другой, а allow_overlap не передан или равен false."],
            ["409", "idempotency_conflict", "Ключ идемпотентности уже использован с другим телом."],
            ["413", "payload_too_large", "Файл истории больше допустимого размера."],
            ["422", "validation_error", "Поля не проходят проверку."],
            ["422", "in_the_past", "Дата или время уже прошли. Детали содержат field."],
            ["422", "invalid_csv", "Файл истории разобрать нельзя. Список ошибок по строкам — в details.errors."],
            ["429", "too_many_requests", "Больше 10 попыток входа в минуту с одного адреса."],
        ],
    )
    s.p("Коды ошибок полей: required, invalid, too_short, too_long, invalid_timezone, invalid_interval, duration_too_long, invalid_range, invalid_type, invalid_status.")
    s.h("Время в прошлом", 2)
    s.code(
        """
{
  "code": "in_the_past",
  "message": "Указанное время уже прошло",
  "details": { "field": "starts_at" }
}
"""
    )
    s.h("Пересечение встреч", 2)
    s.code(
        """
{
  "code": "overlap",
  "message": "На это время уже есть встреча",
  "details": {
    "conflicts": []
  }
}
"""
    )
    s.p("В conflicts лежат полные объекты встреч, с которыми пересекается запрошенный интервал. Новая встреча при этом ответе не создаётся и старая не изменяется.")


def auth(s: Spec) -> None:
    s.h("4. Аутентификация", 1)
    s.p("Пользователь задаётся почтой и паролем. Календарь личный: общих календарей и приглашений в этой версии нет.")
    s.h("Объект пользователя", 2)
    s.code(
        """
{
  "id": "6f1c2c4e-8a1b-4d77-9c2a-1b6e5d0a1111",
  "email": "user@example.com",
  "timezone": "Europe/Moscow",
  "created_at": "2026-10-10T21:00:00+03:00"
}
"""
    )
    s.h("POST /api/v1/auth/register", 2)
    s.p("Создаёт пользователя и сразу выдаёт токены. Токен не требуется.")
    s.table(
        ["Поле", "Обяз.", "Правило"],
        [
            ["email", "да", "Почта, до 254 символов. Сервер обрезает пробелы по краям и приводит к нижнему регистру."],
            ["password", "да", "От 8 до 128 символов. Пробелы по краям сохраняются."],
            ["timezone", "нет", "Пояс IANA. Если поле опущено, сервер ставит Europe/Moscow."],
        ],
    )
    s.p("Ответ 201 — объект сессии. Ошибки: 409 email_taken, 422 validation_error.")
    s.h("POST /api/v1/auth/login", 2)
    s.p("Тело: email и password. Ответ 200 — объект сессии. Ошибки: 401 invalid_credentials, 429 too_many_requests.")
    s.h("Объект сессии", 2)
    s.code(
        """
{
  "access_token": "<jwt>",
  "token_type": "Bearer",
  "expires_in": 900,
  "refresh_token": "<opaque>",
  "refresh_expires_in": 2592000,
  "user": {}
}
"""
    )
    s.p(
        "access token живёт 15 минут (expires_in — 900 секунд) и является JWT. "
        "В его полях есть sub с идентификатором пользователя, iat, exp и type со значением access. "
        "Клиент может прочитать срок из expires_in и не обязан разбирать JWT. "
        "refresh token — случайная строка на 30 суток. Это секрет, его передают только в теле refresh и logout."
    )
    s.h("POST /api/v1/auth/refresh", 2)
    s.p("Тело: refresh_token. Токен доступа не требуется. Ответ 200 — новый объект сессии. Прежний refresh token после успеха больше не действует.")
    s.p("Повтор уже использованного или неизвестного refresh token даёт 401 invalid_refresh. Пользователь входит заново.")
    s.h("POST /api/v1/auth/logout", 2)
    s.p("Тело: refresh_token. Токен доступа не требуется. Ответ 204. Повторный выход с тем же refresh token тоже отвечает 204.")
    s.h("GET /api/v1/auth/me", 2)
    s.p("Ответ 200 — объект пользователя.")
    s.h("PATCH /api/v1/auth/me", 2)
    s.p(
        "Тело содержит timezone. Смена пояса не переписывает уже сохранённые моменты: встреча остаётся в тот же абсолютный миг и начинает отображаться в новом поясе. "
        "Срок задачи остаётся той же календарной датой. Ответ 200 — объект пользователя."
    )


def tasks(s: Spec) -> None:
    s.h("5. Задачи", 1)
    s.p("Задача — действие с результатом. Описание обязательно, срок необязателен. В обзор за период попадают только задачи со сроком.")
    s.h("Объект задачи", 2)
    s.code(
        """
{
  "id": "11111111-1111-4111-8111-111111111111",
  "type": "task",
  "description": "подготовить отчёт",
  "due_on": "2026-10-16",
  "status": "active",
  "completed_at": null,
  "source_id": null,
  "created_at": "2026-10-10T21:05:00+03:00",
  "updated_at": "2026-10-10T21:05:00+03:00"
}
"""
    )
    s.table(
        ["Поле", "Значение"],
        [
            ["description", "Текст от 1 до 2000 символов после обрезки пробелов по краям."],
            ["due_on", "Календарная дата либо null."],
            ["status", "active или completed."],
            ["completed_at", "Момент перехода в completed либо null."],
            ["source_id", "Идентификатор из файла истории либо null."],
        ],
    )
    s.h("POST /api/v1/tasks", 2)
    s.p("Создаёт задачу. Заголовок Idempotency-Key рекомендуется. Тело: description и необязательный due_on.")
    s.p("Ответ 201 и заголовок Location: /api/v1/tasks/{id}. Тело — объект задачи. Срок раньше сегодняшнего дня даёт 422 in_the_past и поле due_on.")
    s.h("GET /api/v1/tasks", 2)
    s.table(
        ["Параметр", "Правило"],
        [
            ["status", "active, completed или all. По умолчанию all."],
            ["due", "any, dated или undated. По умолчанию any."],
            ["from и to", "Обе даты вместе, границы включительно. Одна дата без второй даёт 422 validation_error. Задачи без срока в такой выборке отсутствуют. Сочетание с due=undated даёт 422 validation_error с кодом поля invalid_range."],
            ["limit, offset", "Как в общих соглашениях."],
        ],
    )
    s.p("Порядок: сначала active, затем completed; внутри группы срок по возрастанию, задачи без срока в конце; при равенстве — created_at по возрастанию.")
    s.h("GET /api/v1/tasks/{id}", 2)
    s.p("Ответ 200 — объект задачи. Идентификатор встречи или напоминания на этом пути даёт 404.")
    s.h("PATCH /api/v1/tasks/{id}", 2)
    s.p(
        "Меняет только переданные поля description и due_on. Статус этим методом не меняется. "
        "due_on: null снимает срок. Новый срок в прошлом даёт 422 in_the_past, задача остаётся прежней. Ответ 200 — объект задачи."
    )
    s.h("POST /api/v1/tasks/{id}/complete", 2)
    s.p("Переводит задачу в completed и проставляет completed_at. Ответ 200:")
    s.code(
        """
{
  "task": {},
  "status_changed": true
}
"""
    )
    s.p("Если задача уже выполнена, status_changed равен false, completed_at не меняется, код ответа остаётся 200.")
    s.h("POST /api/v1/tasks/{id}/reopen", 2)
    s.p(
        "Возвращает задачу в active и очищает completed_at. Срок не проверяется и не меняется, даже если он уже прошёл. "
        "Ответ устроен так же, как у complete. Повторный возврат уже активной задачи даёт status_changed: false."
    )
    s.h("DELETE /api/v1/tasks/{id}", 2)
    s.p(
        "Удаляет задачу. Заголовок Idempotency-Key рекомендуется. Первый успешный ответ — 204. "
        "Повтор с тем же ключом снова даёт 204. Повтор без ключа, когда записи уже нет, даёт 404 not_found: для ассистента это сигнал сказать, что записи уже нет."
    )


def meetings(s: Spec) -> None:
    s.h("6. Встречи", 1)
    s.p(
        "Встреча занимает интервал. Название, дата и время начала обязательны. "
        "Если конец не передан, сервер ставит начало плюс один час. Конец должен быть строго позже начала, длительность — не больше 24 часов."
    )
    s.p(
        "Пересечение полуоткрытое: встреча 13:00–14:00 и встреча 14:00–15:00 друг другу не мешают, а 13:00–14:00 и 13:30–14:30 мешают. "
        "Свои собственные интервалы при правке в список конфликтов не входят. Задачи и напоминания с встречами не сравниваются."
    )
    s.h("Объект встречи", 2)
    s.code(
        """
{
  "id": "22222222-2222-4222-8222-222222222222",
  "type": "meeting",
  "title": "встреча с дизайнером",
  "starts_at": "2026-10-11T13:00:00+03:00",
  "ends_at": "2026-10-11T14:00:00+03:00",
  "source_id": null,
  "created_at": "2026-10-10T21:10:00+03:00",
  "updated_at": "2026-10-10T21:10:00+03:00"
}
"""
    )
    s.h("Два шага при пересечении", 2)
    s.p("Параллельные встречи разрешены. Сервер записывает такую встречу после явного согласия клиента.")
    s.bullets(
        [
            "Клиент отправляет POST или PATCH с allow_overlap false либо без этого поля.",
            "Если интервал свободен, сервер сохраняет встречу и возвращает 201 или 200 с пустым overlaps.",
            "Если интервал занят, сервер ничего не меняет и возвращает 409 overlap вместе с conflicts.",
            "Ассистент спрашивает, поставить встречу параллельно или выбрать другое время. Форма на сайте показывает то же предупреждение.",
            "Ответ «параллельно» — новый запрос с allow_overlap true и новым Idempotency-Key. Сервер сохраняет встречу и возвращает актуальный список overlaps.",
            "Ответ «другое время» — сервер больше не вызывается, пока клиент не получил новое время. Следующая попытка снова идёт с allow_overlap false.",
            "Новое время тоже может быть занято. Тогда сервер снова отвечает 409, и диалог повторяется.",
        ]
    )
    s.note(
        "Между отказом и повторным запросом другая встреча может появиться. "
        "Флаг allow_overlap true означает согласие на те пересечения, которые есть в момент второго запроса. Сервер не задаёт вопрос ещё раз."
    )
    s.h("POST /api/v1/meetings", 2)
    s.table(
        ["Поле", "Обяз.", "Правило"],
        [
            ["title", "да", "От 1 до 2000 символов после обрезки пробелов."],
            ["starts_at", "да", "Момент начала. Прошедшее время даёт 422 in_the_past."],
            ["ends_at", "нет", "Если опущено, сервер прибавляет один час к началу."],
            ["allow_overlap", "нет", "По умолчанию false."],
        ],
    )
    s.p("Ответ 201, заголовок Location, тело:")
    s.code(
        """
{
  "meeting": {},
  "overlaps": []
}
"""
    )
    s.h("PATCH /api/v1/meetings/{id}", 2)
    s.p("Меняет только переданные поля title, starts_at, ends_at. Поле allow_overlap управляет проверкой и в объект встречи не записывается.")
    s.bullets(
        [
            "Только новое начало: сервер сохраняет прежнюю длительность и сдвигает конец. Час по умолчанию заново не подставляется.",
            "Только новый конец: начало остаётся, конец должен быть позже начала.",
            "Начало и конец вместе: сервер берёт оба значения.",
            "Проверка пересечения выполняется, когда получившийся интервал отличается от текущего. Правка одного названия проверку не включает, даже если встреча уже стоит параллельно с другой.",
            "Новое начало или новый конец в прошлом даёт 422 in_the_past, встреча не меняется.",
            "При конфликте и allow_overlap false ответ 409, встреча не меняется. При согласии сервер обновляет ту же встречу и не создаёт вторую.",
        ]
    )
    s.p("Ответ 200 имеет ту же форму, что и создание: meeting и overlaps.")
    s.h("GET /api/v1/meetings и GET /api/v1/meetings/{id}", 2)
    s.p(
        "Список принимает from и to парой календарных дат. Окно — от 00:00 даты from до конца даты to в поясе пользователя. "
        "В выборку входит встреча, чей интервал хоть сколько-нибудь пересекается с окном. Порядок — по starts_at, затем по id. "
        "Одна встреча отвечает объектом встречи."
    )
    s.h("DELETE /api/v1/meetings/{id}", 2)
    s.p("Правила те же, что у удаления задачи: подтверждение спрашивает клиент, сервер удаляет сразу. Ответ 204, повтор с тем же ключом — 204, повтор без ключа — 404.")


def reminders(s: Spec) -> None:
    s.h("7. Напоминания", 1)
    s.p(
        "Напоминание — отдельная запись со своим текстом и моментом срабатывания. "
        "Оно не является будильником внутри встречи или задачи. Когда момент наступает, сервер создаёт уведомление внутри приложения."
    )
    s.h("Объект напоминания", 2)
    s.code(
        """
{
  "id": "33333333-3333-4333-8333-333333333333",
  "type": "reminder",
  "text": "купить билеты",
  "fire_at": "2026-10-11T18:00:00+03:00",
  "status": "pending",
  "fired_at": null,
  "source_id": null,
  "created_at": "2026-10-10T21:15:00+03:00",
  "updated_at": "2026-10-10T21:15:00+03:00"
}
"""
    )
    s.p("status принимает pending или fired. text — от 1 до 2000 символов. fire_at обязателен.")
    s.h("POST /api/v1/reminders", 2)
    s.p("Тело: text и fire_at. Прошедший момент даёт 422 in_the_past. Ответ 201, заголовок Location, тело — объект напоминания. Заголовок Idempotency-Key рекомендуется.")
    s.h("PATCH /api/v1/reminders/{id}", 2)
    s.bullets(
        [
            "Можно передать text, fire_at или оба поля.",
            "Новый момент в прошлом даёт 422 in_the_past, запись не меняется.",
            "Пока напоминание pending, перенос меняет момент будущего уведомления. Старый момент уведомления не создаёт.",
            "Если напоминание уже fired и новый момент в будущем, статус возвращается в pending, fired_at очищается. Прежнее уведомление в списке остаётся. В новый момент сервер создаст другое.",
            "Правка только текста у fired-напоминания статус не меняет. Текст уже отправленного уведомления тоже не меняется: он фиксируется в момент срабатывания.",
        ]
    )
    s.p("Ответ 200 — объект напоминания.")
    s.h("Срабатывание", 2)
    s.p(
        "Пока сервер работает, уведомление появляется не позже чем через 15 секунд после fire_at. "
        "Если в этот момент сервер был выключен, уведомление появляется в течение 15 секунд после запуска. "
        "В этой версии работает один экземпляр сервера: второй экземпляр мог бы создать уведомление дважды."
    )
    s.p("Удаление напоминания до срабатывания отменяет уведомление. Уже созданное уведомление после удаления напоминания остаётся, его текст хранится в самом уведомлении.")
    s.h("GET и DELETE", 2)
    s.p(
        "GET /api/v1/reminders принимает status со значениями pending, fired или all (по умолчанию all) и пару from, to с теми же правилами окна, что у встреч. "
        "Порядок — по fire_at, затем по id. GET по идентификатору возвращает объект. DELETE устроен так же, как удаление задачи."
    )


def calendar(s: Spec) -> None:
    s.h("8. Списки и календарь", 1)
    s.h("GET /api/v1/calendar", 2)
    s.p("Сводный обзор за период для фраз «что сегодня» и «что на этой неделе». Границы недели и смысл слова «сегодня» вычисляет клиент.")
    s.table(
        ["Параметр", "Правило"],
        [
            ["from, to", "Обязательная пара календарных дат, границы включительно. Длина окна — не больше 366 дней. Если from позже to, ответ 422 validation_error с кодом поля invalid_range."],
            ["types", "Список через запятую: task, meeting, reminder. Если параметр опущен, сервер возвращает все три типа."],
        ],
    )
    s.p("В окно попадают:")
    s.bullets(
        [
            "встреча, чей интервал пересекается с окном;",
            "напоминание, чей fire_at попадает в окно;",
            "задача, чей due_on попадает в окно.",
        ]
    )
    s.p("Задача без срока в этот ответ не входит. Сработавшие напоминания и выполненные задачи за прошлые дни входят, если их дата внутри окна.")
    s.p(
        "Порядок хронологический. Ключ задачи — 00:00 даты due_on в поясе пользователя, ключ встречи — starts_at, ключ напоминания — fire_at. "
        "При равном ключе порядок такой: встреча, напоминание, задача, затем id по возрастанию."
    )
    s.code(
        """
{
  "from": "2026-10-11",
  "to": "2026-10-11",
  "timezone": "Europe/Moscow",
  "items": []
}
"""
    )
    s.p("Элементы items — те же объекты, что у GET по идентификатору. Пустой период — ответ 200 и пустой items. Это не ошибка.")
    s.p(
        "Фраза «покажи мои задачи» периодом не ограничена и включает задачи без срока. Для неё клиент вызывает GET /api/v1/tasks, а не календарь. "
        "Фраза «какие задачи на завтра» уже задаёт период, поэтому задачи без срока в неё не входят."
    )


def search(s: Spec) -> None:
    s.h("9. Поиск", 1)
    s.h("GET /api/v1/search", 2)
    s.p(
        "Один метод и для строки поиска на сайте, и для инструмента ассистента. "
        "Сервер возвращает список совпадений и сам единственную запись не выбирает. Ноль, один или несколько результатов разбирает клиент."
    )
    s.table(
        ["Параметр", "Правило"],
        [
            ["q", "Обязательная строка от 1 до 200 символов после обрезки пробелов."],
            ["types", "Как у календаря. Если опущен, ищутся все три типа."],
            ["from, to", "Необязательная пара дат. Правила те же, что у календаря: записи без даты в окно не входят."],
            ["limit", "По умолчанию 20, наибольшее значение 50."],
        ],
    )
    s.p("Где ищется текст:")
    s.bullets(
        [
            "у задачи — в description;",
            "у встречи — в title;",
            "у напоминания — в text.",
        ]
    )
    s.p(
        "Сравнение не учитывает регистр. Учитываются часть слова и словоформы русского языка: запрос «договор» находит «отправить договор» и формы этого слова. "
        "В этой версии совпадение строится по тексту записи. Поле match зарезервировано под следующий слой: сейчас оно равно text. "
        "Клиент должен спокойно принимать и другие значения match, когда появится смысловой поиск. Форма ответа при этом сохранится."
    )
    s.code(
        """
{
  "query": "договор",
  "items": [
    {
      "item": {},
      "score": 0.91,
      "match": "text"
    }
  ]
}
"""
    )
    s.p(
        "score лежит от 0 до 1, больше — ближе к запросу. Смысл имеет порядок внутри одного ответа; числа разных запросов между собой не сравниваются. "
        "При равном score выше запись с более поздним временем (начало встречи, момент напоминания или срок задачи). Записи без времени внизу. Пустой результат — 200 и пустой items."
    )


def notifications(s: Spec) -> None:
    s.h("10. Уведомления", 1)
    s.p("Уведомление появляется, когда напоминание переходит в fired. Почта и push в этой версии не отправляются.")
    s.h("Объект уведомления", 2)
    s.code(
        """
{
  "id": "44444444-4444-4444-8444-444444444444",
  "reminder_id": "33333333-3333-4333-8333-333333333333",
  "text": "купить билеты",
  "fire_at": "2026-10-11T18:00:00+03:00",
  "created_at": "2026-10-11T18:00:02+03:00",
  "read_at": null
}
"""
    )
    s.p("text и fire_at копируются в момент срабатывания. Если напоминание позже удалено, reminder_id остаётся, а GET напоминания отвечает 404. Показ строится по text.")
    s.h("GET /api/v1/notifications", 2)
    s.p("Параметр unread равен true или false, по умолчанию false. При true в списке только записи с read_at: null. Порядок — created_at по убыванию. Ответ — обычный список с limit и offset.")
    s.h("POST /api/v1/notifications/{id}/read", 2)
    s.p("Ставит read_at, если его ещё не было. Повтор отвечает 200 и оставляет прежний read_at. Тело ответа — объект уведомления.")
    s.h("POST /api/v1/notifications/read-all", 2)
    s.p("Помечает прочитанными все непрочитанные уведомления пользователя. Ответ 200: { \"updated\": 0 }.")
    s.h("WebSocket /api/v1/ws", 2)
    s.p(
        "Адрес: wss://{host}/api/v1/ws?access_token=<access_token>. "
        "Браузерный WebSocket не умеет заголовок Authorization, поэтому токен передаётся в запросе. Сервер проверяет его так же, как Bearer-токен."
    )
    s.bullets(
        [
            "При недействительном токене сервер закрывает соединение кодом 4401.",
            "Сервер отправляет ping каждые 30 секунд.",
            "О создании уведомления сервер присылает одно текстовое сообщение JSON.",
            "Пропущенные события клиент забирает через GET /api/v1/notifications?unread=true после подключения и после возвращения на экран.",
        ]
    )
    s.code(
        """
{
  "type": "notification.created",
  "notification": {}
}
"""
    )


def history(s: Spec) -> None:
    s.h("11. Импорт и экспорт истории", 1)
    s.p("История — это сами записи, а не журнал изменений. Удалённые записи в экспорт не попадают.")
    s.h("Файл", 2)
    s.p("Кодировка UTF-8, разделитель — запятая, кавычки по правилам CSV. В экспорте в начале файла стоит BOM, чтобы Excel открыл кириллицу.")
    s.p("Обязательные колонки: id, date, time, text, type. Дополнительные колонки, которые сервер пишет и принимает: end_time, status.")
    s.table(
        ["Колонка", "Правило"],
        [
            ["id", "Внешний идентификатор строки, до 128 символов. Внутри пользователя он уникален. Это source_id, а не UUID записи."],
            ["date", "ГГГГ-ММ-ДД либо пусто."],
            ["time", "ЧЧ:ММ либо пусто. Секунды не используются."],
            ["text", "Описание задачи, название встречи или текст напоминания."],
            ["type", "task, meeting, reminder. Также принимаются задача, встреча, напоминание. Регистр не важен."],
            ["end_time", "Только для встречи, ЧЧ:ММ. Если пусто, конец равен началу плюс один час."],
            ["status", "Для задачи: active или completed, также активна и выполнена. Для напоминания: pending или fired, также ожидает и сработало. Пустое значение описано ниже."],
        ],
    )
    s.p("Дата и время в файле понимаются в текущем часовом поясе пользователя. Прошедшие значения при импорте разрешены: это уже случившаяся история.")
    s.table(
        ["Тип", "Дата и время", "Что записывается"],
        [
            ["task", "Дата необязательна. Время игнорируется.", "date становится due_on. Пустая date оставляет срок пустым."],
            ["meeting", "Дата и time обязательны.", "Это начало. end_time — конец того же дня. Конец на следующие сутки через этот файл не передаётся."],
            ["reminder", "Дата и time обязательны.", "Это fire_at."],
        ],
    )
    s.p(
        "Импорт не создаёт уведомлений. Напоминание становится pending только если его момент ещё впереди и status не равен fired. "
        "Во всех остальных случаях напоминание сохраняется как fired. Пустой status у задачи означает active. Пустой status у напоминания с будущим моментом означает pending. "
        "У выполненной импортированной задачи completed_at равен времени импорта: в файле отдельного момента выполнения нет."
    )
    s.p("created_at импортированной записи — это момент импорта. Историческая дата живёт в due_on, starts_at или fire_at.")
    s.h("POST /api/v1/history/import", 2)
    s.p("Запрос multipart/form-data, поле файла называется file. Расширение имени не проверяется, сервер смотрит содержимое. Предел — 2 МБ и 5000 строк данных. Заголовок Idempotency-Key рекомендуется.")
    s.p(
        "Файл применяется целиком. Любая ошибка строки отменяет всю загрузку, частичных записей не остаётся. "
        "Номер row — номер строки в файле, заголовок считается первой строкой, первая запись — второй."
    )
    s.p("Если id уже есть у этого пользователя, строка пропускается и существующая запись не перезаписывается. Повтор одного и того же id внутри файла — это ошибка duplicate_id, файл отклоняется.")
    s.p("Пересечения встреч загрузку не останавливают и диалог не открывают. В ответе сервер перечисляет только те новые встречи, которые с кем-то пересекаются, включая пересечения внутри самого файла.")
    s.p("Успешный разбор отвечает 200:")
    s.code(
        """
{
  "created": 10,
  "skipped": 2,
  "skipped_ids": ["t-1", "m-4"],
  "overlaps": [
    {
      "item": {},
      "conflicts": []
    }
  ]
}
"""
    )
    s.p("Ошибка разбора отвечает 422 invalid_csv:")
    s.code(
        """
{
  "code": "invalid_csv",
  "message": "Файл не загружен",
  "details": {
    "errors": [
      {
        "row": 3,
        "code": "missing_time",
        "message": "Для встречи нужно время начала"
      }
    ]
  }
}
"""
    )
    s.p(
        "Коды строк: missing_column, missing_id, missing_text, missing_date, missing_time, invalid_type, invalid_date, invalid_time, "
        "invalid_interval, invalid_status, duplicate_id, duration_too_long. Если не хватает колонки в заголовке, достаточно одной ошибки с row 1."
    )
    s.p("Пример файла:")
    s.code(
        """
id,date,time,text,type,end_time,status
t-1,2026-10-03,,подготовить отчёт,task,,completed
t-2,,,купить билеты,task,,active
m-1,2026-10-10,13:00,встреча с дизайнером,meeting,14:30,
r-1,2026-10-10,18:00,купить цветы,reminder,,pending
"""
    )
    s.h("GET /api/v1/history/export", 2)
    s.p(
        "Ответ 200, Content-Type: text/csv; charset=utf-8, файл history.csv. "
        "Колонки те же: id, date, time, text, type, end_time, status. "
        "В id сервер пишет source_id, а если его нет — UUID записи. Повторный импорт этого файла пропускает уже существующие строки и не создаёт копий."
    )
    s.p(
        "Дата и время выгружаются в текущем поясе пользователя. У задачи time и end_time пустые. "
        "У встречи time — начало, end_time — конец. У напоминания time — момент срабатывания. "
        "status задачи — active или completed, status напоминания — pending или fired, status встречи пустой."
    )
    s.note(
        "Если конец встречи приходится на следующие сутки, экспорт пишет эти часы в end_time. "
        "Повторный импорт считает такой конец тем же днём, что и начало, поэтому интервал через полночь из файла не восстанавливается. "
        "Встреча внутри одного дня при экспорте и обратном импорте сохраняет интервал."
    )


def scenarios(s: Spec) -> None:
    s.h("12. Сценарии вызовов", 1)
    s.p(
        "В таблице — уже разобранное намерение. Относительное время клиент подставляет сам. "
        "Если обязательного поля ещё нет, клиент задаёт вопрос и сервер не вызывает. "
        "Если поиск вернул несколько записей, клиент показывает выбор и до ответа пользователя ничего не меняет."
    )
    s.table(
        ["Ситуация", "Вызовы"],
        [
            ["«Напомни завтра в 18:00 купить билеты»", "POST /reminders с text и fire_at."],
            ["«Напомни вечером купить цветы», час неясен", "Сервер не вызывается, пока клиент не получил момент."],
            ["«Поставь задачу подготовить отчёт к пятнице»", "POST /tasks с description и due_on."],
            ["«Поставь задачу купить билеты»", "POST /tasks только с description."],
            ["«Встреча с дизайнером сегодня в 13:00»", "POST /meetings, ends_at можно не передавать, allow_overlap false."],
            ["«Встреча завтра с 13:00 до 14:00»", "POST /meetings с starts_at и ends_at."],
            ["Время занято, пользователь согласен на параллельную встречу", "Повтор POST или PATCH с allow_overlap true и новым ключом."],
            ["Пользователь хочет другое время", "Сервер не вызывается до нового времени."],
            ["«Что у меня сегодня», «какие планы на сегодня»", "GET /calendar на сегодняшнюю дату, все типы."],
            ["«Есть ли встречи сегодня»", "GET /calendar на сегодня, types=meeting."],
            ["«Что на этой неделе»", "GET /calendar с понедельника по воскресенье в поясе пользователя."],
            ["Период пуст", "Сервер отвечает 200 и пустым списком. Формулировку «ничего нет» собирает клиент."],
            ["«Покажи мои задачи»", "GET /tasks без дат. В списке есть задачи без срока."],
            ["«Какие задачи на завтра»", "GET /tasks или /calendar с датой завтра и types=task. Задач без срока там нет."],
            ["«Найди задачу про договор»", "GET /search?q=договор&types=task."],
            ["«Когда была встреча с Иваном»", "GET /search?q=Иван&types=meeting."],
            ["Найдена одна запись", "Клиент показывает её. Сервер к этому моменту уже отдал список."],
            ["Найдено несколько", "Клиент перечисляет варианты и ждёт выбор."],
            ["Ничего не найдено", "Пустой items. Клиент сообщает, что подходящих записей нет."],
            ["«Удали задачу про отчёт»", "Сначала поиск. После явного согласия — DELETE. Отказ означает, что DELETE не вызывается."],
            ["«Перенеси её на понедельник»", "Клиент берёт id из своего контекста диалога и делает PATCH due_on или starts_at."],
            ["«Перенеси встречу на 15:00»", "PATCH с новым starts_at. Конец сервер сдвинет сам, длительность сохранится."],
            ["«Отметь задачу выполненной»", "POST /tasks/{id}/complete."],
            ["Задача уже выполнена", "Ответ 200 и status_changed false. Клиент говорит, что статус уже такой."],
            ["«Верни задачу в работу»", "POST /tasks/{id}/reopen."],
            ["Две команды в одной фразе", "Два последовательных вызова. Успешный результат первого сохраняется, если второй ответил ошибкой."],
            ["Ответ сервера потерялся, клиент повторяет создание", "Тот же Idempotency-Key и то же тело. Сервер возвращает прежний результат."],
            ["Связи нет до вызова", "Клиент сообщает, что действие выполнить не удалось. Введённые данные остаются у клиента."],
        ],
    )
    s.h("Контекст диалога", 2)
    s.p(
        "Сервер не хранит последнюю запись разговора и незавершённый вопрос. "
        "Модуль ассистента помнит выбранный id, пока операция не завершилась, не отменена или не сменилась новой самостоятельной командой. "
        "К серверу уходит уже конкретный идентификатор."
    )


def out_of_scope(s: Spec) -> None:
    s.h("13. За границами версии 1.0", 1)
    s.p("Этих возможностей в контракте нет, клиентам не нужно закладывать под них отдельные поля ответа.")
    s.bullets(
        [
            "Разбор речи, синтез речи и языковая модель.",
            "Повторяющиеся задачи и встречи.",
            "Предупреждение о возможном дубликате при создании.",
            "Общие календари, участники и приглашения.",
            "Почта и push-уведомления.",
            "Смысловой поиск без общих слов. Когда он появится, это будет тот же GET /search и новое значение match.",
            "Подписка ICS и выгрузка встреч в Google Calendar. Свои записи при этом останутся источником данных, отдельного способа создать встречу в обход этого API не будет.",
            "Двусторонняя синхронизация с календарём Apple или Google.",
        ]
    )
    s.h("Служебный метод", 2)
    s.p("GET /api/v1/health не требует токен и отвечает 200: { \"status\": \"ok\" }. Он нужен для проверки, что сервер запущен, и к календарю отношения не имеет.")


def changelog(s: Spec) -> None:
    s.h("14. История версий", 1)
    s.table(
        ["Версия", "Дата", "Что зафиксировано"],
        [
            [
                "1.0",
                "10 октября 2026",
                "Первая спецификация: личный календарь, три типа записей, пересечение встреч в два шага, выполнение и возврат задачи, уведомления, импорт и экспорт, общий поиск.",
            ],
        ],
    )


if __name__ == "__main__":
    build()
    print(OUT)
